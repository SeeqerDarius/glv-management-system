from pathlib import Path
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "GLV_Dashboard_and_Business_Overview_Guide.docx"
NAVY, PALE, BORDER = "17365D", "F4F7FB", "D9D9D9"

def shade(cell, fill):
    node = OxmlElement("w:shd"); node.set(qn("w:fill"), fill); cell._tc.get_or_add_tcPr().append(node)

def format_cell(cell, fill=None):
    if fill: shade(cell, fill)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}"); node.set(qn("w:val"), "single"); node.set(qn("w:sz"), "4"); node.set(qn("w:color"), BORDER); borders.append(node)
    tc_pr.append(borders)
    margins = OxmlElement("w:tcMar")
    for name, value in (("top", 100), ("start", 120), ("bottom", 100), ("end", 120)):
        node = OxmlElement(f"w:{name}"); node.set(qn("w:w"), str(value)); node.set(qn("w:type"), "dxa"); margins.append(node)
    tc_pr.append(margins)

def metric_table(doc, rows):
    table = doc.add_table(rows=1, cols=4); table.alignment = WD_TABLE_ALIGNMENT.CENTER; table.autofit = False
    widths = [Inches(1.35), Inches(2.35), Inches(2.25), Inches(1.15)]
    for i, text in enumerate(("Figure", "What it means", "How GLV calculates it", "How to read it")):
        cell = table.rows[0].cells[i]; cell.width = widths[i]; format_cell(cell, NAVY)
        run = cell.paragraphs[0].add_run(text); run.bold = True; run.font.color.rgb = RGBColor(255,255,255); run.font.size = Pt(8.5)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for n, row in enumerate(rows):
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cells[i].width = widths[i]; format_cell(cells[i], PALE if n % 2 else None)
            p = cells[i].paragraphs[0]; p.paragraph_format.space_after = Pt(0)
            run = p.add_run(text); run.font.size = Pt(8.2); run.bold = i == 0
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level); p.paragraph_format.keep_with_next = True

def image(doc, filename, caption):
    path = ROOT / "documentation" / "prepared_screenshots" / filename
    if path.exists():
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(path), width=Inches(6.9))
        p = doc.add_paragraph(caption); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.runs[0].italic = True; p.runs[0].font.size = Pt(8)

def build():
    doc = Document(); section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(.65); section.left_margin = section.right_margin = Inches(.7)
    doc.styles["Normal"].font.name = "Aptos"; doc.styles["Normal"].font.size = Pt(9.5); doc.styles["Normal"].paragraph_format.space_after = Pt(6)
    for name, size in (("Title", 26), ("Heading 1", 17), ("Heading 2", 12)):
        style = doc.styles[name]; style.font.name = "Aptos Display"; style.font.size = Pt(size); style.font.color.rgb = RGBColor(0,0,0); style.font.bold = True
    header = section.header.paragraphs[0]; header.text = "ROCK FROST GROUP  |  GLV MANAGEMENT SYSTEM"; header.alignment = WD_ALIGN_PARAGRAPH.RIGHT; header.runs[0].font.size = Pt(8); header.runs[0].font.bold = True
    footer = section.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER; footer.add_run("GLV Dashboard Figures Guide   •   Internal management reference"); footer.runs[0].font.size = Pt(8)

    doc.add_paragraph("GLV Dashboard and Business Overview Figures Guide", style="Title")
    p = doc.add_paragraph("A practical reference for understanding the figures shown to administrators and staff"); p.runs[0].font.size = Pt(13); p.runs[0].font.bold = True
    doc.add_paragraph("Prepared for God’s Love Ventures management and operations teams by Rock Frost Group")
    image(doc, "01-dashboard.png", "Administrator dashboard overview")
    doc.add_paragraph("This guide explains what each figure represents, the records included in its calculation, and the management decision it supports. Dashboard values summarize records entered in GLV. If a payment, account status, product cost, salary, deposit, or refund is missing or incorrect, the related figure will also be incomplete or incorrect.")
    heading(doc, "How to use the figures")
    doc.add_paragraph("Use counts to monitor workload and account health. Use collection and receivable figures to understand customer cash flow. Use product cost, procurement, payroll, and profit figures together; no single card represents cash in the bank or final accounting profit. Currency figures use the system’s configured currency, normally Ghana cedis.")
    doc.add_paragraph("Account cost means product cost price plus transport cost. Included accounts exclude cancelled, closed, and archived accounts unless a figure explicitly says otherwise. The salary due month is the previous calendar month, while current month payroll uses salaries effective in the current month.")

    heading(doc, "Administrator dashboard figures")
    metric_table(doc, [
        ("Total Customers", "All customer profiles in GLV.", "Count of customer records, whether or not each customer has an active plan.", "Size of the customer register."),
        ("Total Staff", "All staff profiles in GLV.", "Count of staff records, including active and inactive profiles.", "Recorded workforce, not only active staff."),
        ("Active Accounts", "Plans currently collecting normally.", "Count of accounts whose effective lifecycle status is Active.", "Current collection workload."),
        ("Completed and Delivered", "Plans fully completed and marked delivered.", "Delivered accounts with Completed or Archived status.", "Confirms fulfilment, not only payment completion."),
        ("Overdue Accounts", "Plans past expected end date with balance remaining.", "Count of accounts whose effective status is Overdue.", "Requires collection follow-up."),
        ("Open Credits or Refunds", "Unused customer credit still owed or available.", "Sum of remaining amounts on Open customer-credit records.", "Customer value still to resolve."),
        ("Payments Collected", "All recorded customer payments.", "Sum of every Payment amount in the system.", "Cumulative, not current-month collections."),
        ("Expected Receivables", "Balance expected from live plans.", "Balances on Active and Overdue accounts.", "Future collection focus."),
        ("Profit Estimate", "Expected gross profit across included plans.", "Sum of target amount minus product and transport cost for each included account.", "Projection before payroll and other overheads."),
    ])
    heading(doc, "Dashboard gain and loss summary")
    metric_table(doc, [
        ("Total Product Exposure", "Capital cost committed to included plans.", "Product cost plus transport cost for accounts not Cancelled, Closed, or Archived.", "Capital tied to plans."),
        ("Procurement Due Now", "Estimated cost of products ready to procure.", "Total procurement-list cost for qualifying accounts.", "Cash needed for current fulfilment."),
        ("Cash After Procurement", "Collections left after procurement due.", "All payments collected minus Procurement Due Now.", "Negative means procurement is not covered."),
        ("Total Expected Profit", "Projected gross profit on included plans.", "Targets minus product and transport costs.", "Compare with payroll and overheads."),
        ("Salary Paid for Due Month", "Payroll recorded against the previous salary month.", "Salary payments classified by salary month.", "Payment date can be in another month."),
        ("Current Month Payroll", "Salary commitment for active staff this month.", "Effective current-month salaries for active staff.", "Expected current payroll obligation."),
        ("Outstanding Salaries", "Unpaid previous-month payroll.", "Due-month payroll minus paid due-month salary, never below zero.", "Zero means recorded due payroll is covered."),
        ("Payroll vs Income", "Collections remaining after due-month salary payments.", "All payments collected minus due-month salary paid.", "Headroom, not net profit."),
        ("Operating Cash Position", "Collections after procurement due and due-month payroll.", "Collections minus procurement due minus due-month salary paid.", "Negative means a recorded cash deficit."),
        ("Projected Net Profit", "Expected profit after current-month payroll.", "Total Expected Profit minus Current Month Payroll.", "Excludes other unrecorded overheads."),
        ("Current Position", "Text status for operating cash.", "Cash Positive at zero or above; otherwise Cash Deficit.", "Quick present-position label."),
        ("Projection", "Text status for projected net profit.", "Projected Profit above zero, Projected Loss below zero, Break Even at zero.", "Forward-looking label."),
        ("Refund Items", "Number of unresolved credit records.", "Count of Open customer credits.", "Review until applied or refunded."),
        ("Closure Refunds", "Open refunds created by account closure.", "Open credits sourced from Account Closure Refund.", "Closed-plan liabilities needing attention."),
    ])
    heading(doc, "Staff dashboard figures")
    doc.add_paragraph("Staff users see only records assigned to their staff profile. Their figures do not represent the whole business.")
    metric_table(doc, [
        ("My Customers", "Customers assigned to the signed-in staff member.", "Count of distinct assigned customer profiles.", "The staff member’s portfolio."),
        ("My Accounts", "Plans belonging to assigned customers.", "Count of assigned customer accounts across statuses.", "One customer may have several accounts."),
        ("Active Accounts", "Assigned plans currently active.", "Count with effective Active status.", "Immediate collection workload."),
        ("Payments Today", "Payment entries recorded today.", "Count of today’s qualifying payment records.", "Transaction count, not value."),
        ("Collected Today", "Value collected today.", "Sum of today’s qualifying payment amounts.", "Compare with receipts and targets."),
        ("Collected This Week", "Value collected this week.", "Sum of qualifying payments in the current week.", "Weekly collection performance."),
    ])

    heading(doc, "Reports Business Overview")
    image(doc, "10-reports.png", "Reports page showing Business Overview and weekly reporting controls")
    doc.add_paragraph("The Reports page combines lifetime, selected-week, current-month, and due-month measures. Read the period stated below before comparing cards.")
    doc.add_page_break()
    metric_table(doc, [
        ("Total Collected", "All customer payments recorded.", "Sum of all Payment amounts.", "Lifetime recorded collections."),
        ("Recorded This Week", "Payments credited to staff activity for the selected week.", "Weekly collection amounts across staff rows.", "Compare with weekly deposits."),
        ("Deposited This Week", "Staff deposits entered for the selected week.", "Sum of Staff Deposit amounts dated in that week.", "Deposits, not customer payment entries."),
        ("Weekly Deposit Variance", "Difference between deposits and collections.", "Deposited This Week minus Recorded This Week.", "Negative shortage; positive surplus; zero balanced."),
        ("Outstanding Balance", "Balance on all included accounts.", "Balances excluding Cancelled, Closed, and Archived accounts.", "Broader than Expected Receivables."),
        ("Expected Receivables", "Balance on collectible live plans.", "Balances on Active and Overdue accounts.", "Expected future collections."),
        ("Product Cost Exposure", "Product and transport capital across included plans.", "Product cost plus transport cost, excluding Cancelled, Closed, and Archived accounts.", "Recorded capital exposure."),
        ("Current Month Payroll", "Salary commitment for active staff this month.", "Effective current-month salaries for active staff.", "Paid and unpaid obligation."),
        ("Salary Paid for Due Month", "Salary allocated to the previous month.", "Sum by salary month, not payment date.", "Use with Outstanding Salaries."),
        ("Outstanding Salaries", "Previous-month payroll still unpaid.", "Due payroll minus paid due-month salary, floored at zero.", "Payroll arrears."),
        ("Payroll vs Income", "Current-month collections after due-month salary paid.", "Current-month payment income minus due-month salary paid.", "Negative means payroll exceeds compared income."),
        ("Payroll % of Revenue", "Share of current-month income used for due-month salary.", "Due-month salary paid divided by current-month income times 100; zero if income is zero.", "Higher leaves less collection headroom."),
        ("Net Profit So Far", "Collections after included product exposure and due-month salary.", "Total Collected minus Product Cost Exposure minus Salary Paid for Due Month.", "Management estimate, not statutory profit."),
        ("Projected Profit Loss or Break Even", "Expected profit after current-month payroll.", "Total Expected Profit minus Current Month Payroll; label follows its sign.", "Forward-looking result from current records."),
    ])
    heading(doc, "Common reasons a figure changes")
    for text in (
        "Recording, editing, or deleting a customer payment changes collection, balance, receivable, variance, and profit-related views.",
        "Changing account lifecycle or delivery status changes counts and may include or exclude its balance and cost.",
        "Updating product cost price or transport cost changes exposure and expected profit.",
        "Recording salary against a salary month changes due-month figures even when payment happens in another month.",
        "Recording staff deposits changes Deposited This Week and weekly variance; it does not create a customer payment.",
        "Changing an open credit to applied or refunded changes refund cards without changing payment history.",
    ): doc.add_paragraph(text, style="List Bullet")
    heading(doc, "Checks before relying on the figures")
    doc.add_paragraph("Confirm that payment dates and amounts match receipts, deposits use the correct week, salary payments use the correct salary month, product and transport costs are current, delivered products are marked delivered, and closed or cancelled accounts have been processed correctly. Investigate weekly shortages rather than carrying them forward without explanation.")
    doc.add_paragraph("These figures support daily decisions. Formal financial statements still require complete expense records, bank reconciliation, inventory valuation, tax treatment, and accounting review beyond the operational calculations described here.")
    doc.core_properties.title = "GLV Dashboard and Business Overview Figures Guide"; doc.core_properties.subject = "Definitions and calculations for GLV management figures"; doc.core_properties.author = "Rock Frost Group"
    doc.save(OUT); print(OUT)

if __name__ == "__main__": build()
